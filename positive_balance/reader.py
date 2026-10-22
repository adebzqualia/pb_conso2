"""Read unchanged country templates and FX tables into auditable snapshots.

Formula results must have been calculated and saved in Excel by the sender.
This reader consumes caches and never recalculates or writes source workbooks.
"""
from __future__ import annotations

import fnmatch
import hashlib
import json
import math
import re
import warnings
from datetime import date, datetime
from pathlib import Path

import openpyxl
from openpyxl.utils import column_index_from_string, get_column_letter

from .model import complete_sum, parse_period, quarter_end, rolling_months

CATEGORIES = ("refund", "unreachable", "false_positive", "authority_refund", "other_without_refund")
COUNT_FIELDS = {"count", "processed_count", "remaining_count_source"}


def _norm(value):
    return " ".join(str(value).split()).casefold()


class Book:
    def __init__(self, path, issues, tokens):
        self.path, self.issues, self.tokens = Path(path), issues, tokens
        with warnings.catch_warnings():
            # The source print area is irrelevant to read-only extraction.
            warnings.filterwarnings("ignore", message="Print area cannot be set to Defined name:.*", category=UserWarning)
            self.formulas = openpyxl.load_workbook(path, data_only=False, read_only=False)
            self.values = openpyxl.load_workbook(path, data_only=True, read_only=False)

    def close(self):
        self.formulas.close()
        self.values.close()

    def source(self, sheet, cell):
        return f"{self.path.name}!{sheet}!{cell}"

    def value(self, sheet, cell):
        if sheet not in self.values:
            raise ValueError(f"{self.path.name}: missing sheet {sheet}")
        formula, cached = self.formulas[sheet][cell], self.values[sheet][cell]
        if formula.data_type == "f" and cached.value is None:
            raise ValueError(f"Missing formula cache: {self.source(sheet, cell)}. Recalculate and save in Excel.")
        if cached.data_type == "e":
            raise ValueError(f"Excel error {cached.value}: {self.source(sheet, cell)}")
        return cached.value

    def number(self, sheet, cell, entity, factor=1, count=False):
        value = self.value(sheet, cell)
        source = self.source(sheet, cell)
        if value is None or (isinstance(value, str) and value.strip() in self.tokens):
            self.issues.append(issue("warning", "MISSING_VALUE", entity, "Numeric source value is missing; retained as unknown.", source))
            return None
        if isinstance(value, bool) or not isinstance(value, (int, float)) or not math.isfinite(value):
            raise ValueError(f"Expected finite number at {source}, got {value!r}")
        if count and value != int(value):
            raise ValueError(f"Non-integral contract count at {source}: {value}")
        return int(value) if count else value * factor

    def markers(self, sheet, markers):
        for cell, expected in markers.items():
            if _norm(expected) not in _norm(self.value(sheet, cell)):
                raise ValueError(f"Unexpected layout at {self.source(sheet, cell)}; expected label containing {expected!r}")


def issue(severity, code, entity, message, source=""):
    return dict(severity=severity, code=code, entity=entity, message=message, source=source)


def _date(value, source):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    raise ValueError(f"Expected Excel date at {source}, got {value!r}")


def _vintage(value):
    if isinstance(value, (int, float)) and value == int(value):
        return str(int(value))
    value = str(value).strip()
    if not re.fullmatch(r"<?\s*20\d{2}", value):
        raise ValueError(f"Invalid vintage label {value!r}")
    return value.replace(" ", "")


def _validate_vintages(vintages, period, entity, location):
    if not vintages or not vintages[0].startswith("<"):
        raise ValueError(f"{entity}: {location} vintages must start with an older-year bucket")
    first_year = int(vintages[0][1:])
    expected = [f"<{first_year}"] + [str(year) for year in range(first_year, int(period[:4]) + 1)]
    if vintages != expected:
        raise ValueError(f"{entity}: non-contiguous or wrong-period {location} vintages {vintages}; expected {expected}")


def _check(issues, entity, name, expected, actual, source, tolerance):
    if expected is not None and actual is not None and abs(expected - actual) > tolerance:
        issues.append(issue("warning", "RECONCILIATION", entity,
                            f"{name}: detailed/calculated {expected:.8g} versus reported {actual:.8g}.", source))


def _fields(book, sheet, column, mapping, entity, factor):
    record = {key: book.number(sheet, f"{column}{row}", entity, factor,
                               key in COUNT_FIELDS) for key, row in mapping.items()}
    # None here means not supplied by this profile, not an invented zero.
    for key in CATEGORIES:
        record.setdefault(key, None)
    record["source"] = book.source(sheet, f"{column}{min(mapping.values())}:{column}{max(mapping.values())}")
    return record


def _entity(book, spec, profile, kpi3, period, issues, tolerance):
    eid, factor = spec["id"], spec.get("units", 1)
    if isinstance(factor, bool) or not isinstance(factor, (int, float)) or not math.isfinite(factor) or factor <= 0:
        raise ValueError(f"{eid}: units must be a finite positive amount multiplier")
    prefix = spec.get("sheet_prefix", "")
    s1, s2, s3 = (f"{prefix}KPI_{n}" for n in (1, 2, 3))
    for sheet in (s1, s2, s3):
        if sheet not in book.values:
            raise ValueError(f"{eid}: missing required sheet {sheet}")
    book.markers(s2, profile["markers"])
    book.markers(s3, kpi3["markers"])
    for sheet in (s1, s2):
        reported_country = book.value(sheet, profile["country_cell"])
        if _norm(reported_country) not in {_norm(v) for v in spec.get("country_aliases", [spec["name"]])}:
            raise ValueError(f"{eid}: country mismatch at {book.source(sheet, profile['country_cell'])}: {reported_country!r}")
        currency = book.value(sheet, profile["currency_cell"])
        if _norm(currency) not in {_norm(v) for v in spec.get("currency_aliases", [spec["currency"]])}:
            raise ValueError(f"{eid}: currency {currency!r} disagrees with configured {spec['currency']}")
        cell = profile["period_cell"]
        if _date(book.value(sheet, cell), book.source(sheet, cell)) != quarter_end(period):
            raise ValueError(f"{eid}: reporting date in {sheet}!{cell} does not match {period} quarter end")
    if _date(book.value(s3, kpi3["period_cell"]), book.source(s3, kpi3["period_cell"])) != quarter_end(period):
        raise ValueError(f"{eid}: KPI_3 reporting date does not match {period}")
    entity = {key: spec[key] for key in ("id", "name", "currency", "profile", "file")}
    entity.update(report_date=quarter_end(period).isoformat(), units=factor,
                  processed_fields=profile["processed_fields"], requires_cst=spec.get("requires_cst", False),
                  stock=[], flows=[], stock_bins=[], stock_vintages=[])
    smap = profile["stock"]
    columns = range(column_index_from_string(smap["first_column"]), column_index_from_string(smap["last_column"]) + 1)
    for index in columns:
        col = get_column_letter(index)
        record = _fields(book, s2, col, smap["fields"], eid, factor)
        record["vintage"] = _vintage(book.value(s2, f"{col}{smap['header_row']}"))
        entity["stock"].append(record)
    vintages = [r["vintage"] for r in entity["stock"]]
    _validate_vintages(vintages, period, eid, "KPI_2")
    entity["stock_totals"] = _fields(book, s2, smap["total_column"], smap["fields"], eid, factor)
    for field in smap["fields"]:
        _check(issues, eid, f"KPI_2 vintage {field} total", complete_sum(r[field] for r in entity["stock"]),
               entity["stock_totals"][field], entity["stock_totals"]["source"], tolerance)
    fmap = profile["flows"]
    months = []
    for col in fmap["columns"]:
        cell = f"{col}{fmap['header_row']}"
        month = _date(book.value(s2, cell), book.source(s2, cell)).strftime("%Y-%m")
        months.append(month)
        record = _fields(book, s2, col, fmap["fields"], eid, factor)
        record["month"] = month
        entity["flows"].append(record)
    if months != rolling_months(period, len(fmap["columns"])):
        raise ValueError(f"{eid}: rolling month headers {months} do not match {period}")
    entity["flows"].sort(key=lambda record: record["month"])
    def stock_row(row, label_key):
        return {label_key: str(book.value(s3, f"{kpi3['label_column']}{row}")).strip(),
                "count": book.number(s3, f"{kpi3['count_column']}{row}", eid, count=True),
                "amount": book.number(s3, f"{kpi3['amount_column']}{row}", eid, factor),
                "source": book.source(s3, f"{kpi3['label_column']}{row}:{kpi3['amount_column']}{row}")}
    for row in range(kpi3["bins_first_row"], kpi3["bins_last_row"] + 1):
        if not str(book.value(s3, f"{kpi3['label_column']}{row}") or "").strip():
            raise ValueError(f"{eid}: missing KPI_3 amount-band label at row {row}")
        entity["stock_bins"].append(stock_row(row, "band"))
    entity["stock_bins_totals"] = stock_row(kpi3["bins_total_row"], "label")
    if not _norm(entity["stock_bins_totals"]["label"]).startswith("total"):
        raise ValueError(f"{eid}: expected KPI_3 amount-band total label")
    for row in range(kpi3["vintages_first_row"], kpi3["vintages_max_row"] + 1):
        label = book.value(s3, f"{kpi3['label_column']}{row}")
        if _norm(label).startswith("total"):
            entity["stock_vintages_totals"] = stock_row(row, "label")
            break
        record = stock_row(row, "vintage")
        record["vintage"] = _vintage(label)
        entity["stock_vintages"].append(record)
    else:
        raise ValueError(f"{eid}: no KPI_3 vintage total row found")
    vs = [r["vintage"] for r in entity["stock_vintages"]]
    _validate_vintages(vs, period, eid, "KPI_3")
    for collection, total_key in (("stock_bins", "stock_bins_totals"), ("stock_vintages", "stock_vintages_totals")):
        for field in ("count", "amount"):
            _check(issues, eid, f"{collection} {field} total", complete_sum(r[field] for r in entity[collection]),
                   entity[total_key][field], entity[total_key]["source"], 0 if field == "count" else tolerance)
    _check(issues, eid, "KPI_3 bins / vintage total", entity["stock_bins_totals"]["amount"],
           entity["stock_vintages_totals"]["amount"], entity["stock_vintages_totals"]["source"], tolerance)
    _check(issues, eid, "KPI_2 / KPI_3 remaining", entity["stock_totals"]["remaining_source"],
           entity["stock_vintages_totals"]["amount"], entity["stock_vintages_totals"]["source"], tolerance)
    for record in entity["stock"] + entity["flows"] + [entity["stock_totals"]]:
        processed = complete_sum(record[key] for key in profile["processed_fields"])
        _check(issues, eid, "Processed component sum", processed,
               record.get("processed_source", record.get("processed_amount")), record["source"], tolerance)
        total = record.get("total", record.get("amount"))
        expected_remaining = None if total is None or processed is None else total - processed
        _check(issues, eid, "Remaining amount", expected_remaining,
               record.get("remaining_source", record.get("remaining_amount_source")), record["source"], tolerance)
        if "count" in record:
            expected_count = None if record["count"] is None or record["processed_count"] is None else record["count"] - record["processed_count"]
            _check(issues, eid, "Remaining contract count", expected_count,
                   record["remaining_count_source"], record["source"], 0)
    if entity["currency"] != "EUR":
        label_cell = f"A{fmap['fields']['amount']}"
        if re.search(r"\bEUR\b|€", str(book.value(s2, label_cell)), re.IGNORECASE):
            issues.append(issue("warning", "CURRENCY_LABEL_MISMATCH", eid,
                                f"Amount label says EUR, but validated currency metadata is {entity['currency']}; metadata/configuration controls conversion.",
                                book.source(s2, label_cell)))
    bins = entity["stock_bins"]
    if bins and bins[0]["count"] is not None and bins[0]["count"] > 0 and bins[0]["amount"] is not None and bins[0]["amount"] > 0:
        if all(item["count"] == 0 and item["amount"] == 0 for item in bins[1:]):
            issues.append(issue("warning", "BUCKET_CONCENTRATION", eid,
                                "All reported contracts and amounts are in the first amount band; verify that a total was not entered in place of the distribution.", bins[0]["source"]))
    if entity["requires_cst"]:
        issues.append(issue("warning", "CST_REQUIRED", eid, "Separate CST balance is not supplied in the country workbook; complete stock including CST is unavailable.", book.path.name))
    return entity


def extract_folder(folder: Path, config_dir: Path, period: str | None = None) -> dict:
    """Extract configured entities, reject structural/period/FX errors, preserve unknowns."""
    folder, config_dir = Path(folder), Path(config_dir)
    if not folder.is_dir():
        raise ValueError(f"Input folder does not exist: {folder}")
    layouts = json.loads((config_dir / "layouts.json").read_text(encoding="utf-8"))
    countries = json.loads((config_dir / "countries.json").read_text(encoding="utf-8"))
    if layouts.get("schema_version") != 1 or countries.get("schema_version") != 1:
        raise ValueError("Unsupported configuration schema_version; expected 1")
    if layouts["fx"].get("rate_type") != "closing" or layouts["fx"].get("quote") != "local_per_EUR":
        raise ValueError("Base reader supports closing FX rates quoted in local currency units per EUR only")
    for profile_name, profile in layouts["profiles"].items():
        fields = profile.get("processed_fields", [])
        if not fields or len(set(fields)) != len(fields) or any(key not in CATEGORIES for key in fields):
            raise ValueError(f"{profile_name}: processed_fields must contain unique, supported category names")
        if any(key not in profile[section]["fields"] for section in ("stock", "flows") for key in fields):
            raise ValueError(f"{profile_name}: every processed category requires stock and flow cell mappings")
        columns = profile["flows"]["columns"]
        if len(columns) != 6 or len(set(columns)) != 6:
            raise ValueError(f"{profile_name}: base reader requires six distinct rolling-month columns")
    target = parse_period(period or folder.name)
    # An explicit CLI period must not disguise a contradictory period-bearing folder.
    try:
        folder_period = parse_period(folder.name)
    except ValueError:
        folder_period = target
    if folder_period != target:
        raise ValueError(f"Folder period {folder_period} disagrees with requested {target}")
    result = dict(schema_version=1, period=target, entities=[], fx=[], issues=[], sources=[], excluded=[])
    result["configurations"] = [{"file": filename, "sha256": hashlib.sha256((config_dir / filename).read_bytes()).hexdigest()}
                                for filename in ("layouts.json", "countries.json")]
    issues, books, files = result["issues"], {}, {}
    for path in folder.iterdir():
        if path.is_file():
            if path.name.casefold() in files:
                raise ValueError(f"Duplicate filename ignoring case: {path.name}")
            files[path.name.casefold()] = path
    specs = countries["entities"]
    for spec in specs:
        if not re.fullmatch(r"[A-Za-z0-9_-]+", spec["id"]):
            raise ValueError("Entity IDs must contain only letters, digits, underscores or hyphens")
        if not re.fullmatch(r"[A-Z]{3}", spec["currency"]):
            raise ValueError(f"{spec['id']}: currency must be a three-letter uppercase code")
    if len({s["id"] for s in specs}) != len(specs):
        raise ValueError("Duplicate entity IDs in countries.json")
    identities = [(s["file"].casefold(), s.get("sheet_prefix", "")) for s in specs if s.get("enabled", True)]
    if len(set(identities)) != len(identities):
        raise ValueError("Duplicate workbook/sheet mappings would double-count an entity")
    known = {s["file"].casefold() for s in specs} | {layouts["fx"]["file"].casefold()}
    for exclusion in countries.get("excluded_files", []):
        known.add(exclusion["file"].casefold())
        if exclusion["file"].casefold() in files:
            result["excluded"].append(exclusion)
            issues.append(issue("warning", "EXCLUDED_FILE", "", exclusion["reason"], exclusion["file"]))
    for name, path in files.items():
        ignored = any(fnmatch.fnmatch(name, pattern.casefold()) for pattern in layouts.get("ignore_file_patterns", []))
        if path.suffix.lower() in {".xlsx", ".xlsm", ".xls", ".xlsb"} and name not in known and not ignored:
            raise ValueError(f"Unconfigured Excel input {path.name}; add an entity, explicit exclusion, or ignore pattern")
    def get_book(filename):
        key = filename.casefold()
        if key not in files:
            raise ValueError(f"Missing input file {filename}")
        if key not in books:
            books[key] = Book(files[key], issues, set(layouts.get("missing_numeric_tokens", ["", "-"])))
            result["sources"].append(dict(file=files[key].name, sha256=hashlib.sha256(files[key].read_bytes()).hexdigest()))
        return books[key]
    try:
        for spec in specs:
            if not spec.get("enabled", True):
                if spec["file"].casefold() in files:
                    result["excluded"].append({"file": spec["file"], "entity": spec["id"], "reason": "Disabled in configuration"})
                continue
            if spec["file"].casefold() not in files and not spec.get("required", True):
                issues.append(issue("warning", "OPTIONAL_FILE_MISSING", spec["id"], "Configured optional country input is absent.", spec["file"]))
                continue
            if spec["profile"] not in layouts["profiles"]:
                raise ValueError(f"Unknown layout profile {spec['profile']}")
            result["entities"].append(_entity(get_book(spec["file"]), spec, layouts["profiles"][spec["profile"]],
                                               layouts["kpi3"], target, issues, layouts.get("amount_tolerance", .01)))
        fx = layouts["fx"]
        book, sheet = get_book(fx["file"]), fx["sheet"]
        book.markers(sheet, {fx["header_cell"]: fx["header_contains"]})
        fx_date = _date(book.value(sheet, fx["date_cell"]), book.source(sheet, fx["date_cell"]))
        if (fx_date.year, fx_date.month) != (quarter_end(target).year, quarter_end(target).month):
            raise ValueError(f"FX date {fx_date} does not match {target} closing month")
        rates = {}
        for row in range(fx["first_row"], book.values[sheet].max_row + 1):
            code = book.value(sheet, f"{fx['currency_column']}{row}")
            if code is None:
                continue
            code = str(code).strip().upper()
            if code in rates:
                raise ValueError(f"Duplicate FX currency {code}")
            cell = f"{fx['rate_column']}{row}"
            rate = book.number(sheet, cell, "FX")
            if rate is None or rate <= 0:
                raise ValueError(f"Missing or nonpositive FX rate for {code}")
            rates[code] = dict(currency=code, rate=rate, source=book.source(sheet, cell),
                               rate_type=fx.get("rate_type", "closing"), quote=fx.get("quote", "local_per_EUR"), date=fx_date.isoformat())
        if "EUR" not in rates or abs(rates["EUR"]["rate"] - 1) > 1e-12:
            raise ValueError("EUR closing FX rate must be 1 for a local_per_EUR table")
        for currency in sorted({entity["currency"] for entity in result["entities"]}):
            if currency not in rates:
                raise ValueError(f"Missing FX currency {currency}")
            result["fx"].append(rates[currency])
        if not result["entities"]:
            raise ValueError("No enabled country data found")
    finally:
        for book in books.values():
            book.close()
    return result
