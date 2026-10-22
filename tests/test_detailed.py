"""Presentation-report regressions against source controls and native Excel charts.

The reference slides define the layout, not replacement numbers. In particular,
the country KPI_3 distribution and the global stock including unreachable clients
have different scopes, and M-5 efficiency is not semester-weighted efficiency.
"""
from __future__ import annotations

from contextlib import redirect_stderr, redirect_stdout
from copy import deepcopy
import hashlib
import io
import json
from pathlib import Path
import tempfile
import unittest
import xml.etree.ElementTree as ET
from zipfile import ZipFile

import openpyxl

from positive_balance.detail_data import build_presentation_data, previous_period
from positive_balance.detailed_workbook import write_detailed_workbook
from positive_balance.history import merge_snapshot
from positive_balance.model import rolling_months
from positive_balance.reader import extract_folder


ROOT = Path(__file__).resolve().parents[1]
CHART_NS = {"c": "http://schemas.openxmlformats.org/drawingml/2006/chart"}


def country(data, identifier):
    return next(item for item in data["countries"] if item["id"] == identifier)


def source_entity(snapshot, identifier):
    return next(item for item in snapshot["entities"] if item["id"] == identifier)


def period_copy(snapshot, period, rate_multiplier=1):
    """Synthetic prior observation with the same local values but different FX."""
    result = deepcopy(snapshot)
    result["period"] = period
    for entity in result["entities"]:
        for record, month in zip(entity["flows"], reversed(rolling_months(period))):
            record["month"] = month
    for rate in result["fx"]:
        if rate["currency"] != "EUR":
            rate["rate"] *= rate_multiplier
    return result


class DetailedDataTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.snapshot = extract_folder(ROOT / "Q2_2026", ROOT / "config")
        cls.history = merge_snapshot({"snapshots": {}}, cls.snapshot)
        cls.config = json.loads((ROOT / "config/presentation.json").read_text(encoding="utf-8"))
        cls.expected = json.loads((ROOT / "analysis/reconciliation_expected.json").read_text(encoding="utf-8"))
        cls.data = build_presentation_data(cls.history, "2026Q2", cls.config)

    def assertMoney(self, actual, expected):
        self.assertIsNotNone(actual)
        self.assertAlmostEqual(actual, expected, delta=.00001)

    def test_country_local_currency_and_euro_summary_reconcile_to_inputs(self):
        self.assertEqual({c["id"] for c in self.data["countries"]}, {"AT", "IT", "ZA", "DK", "NO", "SE"})
        self.assertEqual(self.data["months"], [f"2026-{month:02d}" for month in range(1, 7)])
        for item in self.data["countries"]:
            expected = self.expected[item["name"]]
            with self.subTest(country=item["id"]):
                self.assertEqual(item["currency"], expected["currency"])
                self.assertMoney(item["metrics"]["remaining_eur"], expected["remaining_eur"])
                self.assertMoney(item["metrics"]["identified_eur"], expected["credit_stock_including_unreachable_eur"])
                self.assertMoney(item["metrics"]["flow_eur"], expected["flow_6m_eur"])
                self.assertMoney(item["metrics"]["flow_local"], expected["flow_6m_local"])
                self.assertMoney(item["metrics"]["monthly_average_local"], expected["flow_6m_local"] / 6)
                self.assertMoney(item["metrics"]["average_balance_local"], expected["flow_6m_local"] / expected["flow_contracts_6m"])
                self.assertMoney(item["metrics"]["oldest_amount"], expected["oldest_cohort_efficiency_amount"])
                self.assertMoney(item["metrics"]["oldest_count"], expected["oldest_cohort_efficiency_count"])

    def test_monthly_stacks_reconcile_without_reclassifying_nordics(self):
        for item in self.data["countries"]:
            entity = source_entity(self.snapshot, item["id"])
            rates = self.expected[item["name"]]["closing_local_per_eur"]
            self.assertEqual([row["month"] for row in item["monthly"]], self.data["months"])
            for month in item["monthly"]:
                source = next(row for row in entity["flows"] if row["month"] == month["month"])
                with self.subTest(country=item["id"], month=month["month"]):
                    treated = sum(source[field] for field in entity["processed_fields"])
                    self.assertMoney(month["processed_amount"], treated)
                    self.assertMoney(month["remaining"] + treated, source["amount"])
                    self.assertMoney(month["amount_eur"], source["amount"] / rates)
                    self.assertMoney(month["efficiency_amount"], treated / source["amount"])
                    self.assertMoney(month["efficiency_count"], source["processed_count"] / source["count"])
                    for field in entity["processed_fields"]:
                        self.assertMoney(month[field], source[field])
        # The extended Nordic source's "without refund" is not a false positive.
        denmark = country(self.data, "DK")
        self.assertGreater(sum(row["other_without_refund"] for row in denmark["monthly"]), 0)
        self.assertTrue(all(row.get("false_positive") is None for row in denmark["monthly"]))

    def test_kpi3_distributions_use_remaining_stock_and_preserve_missing_rows(self):
        for item in self.data["countries"]:
            expected = self.expected[item["name"]]
            with self.subTest(country=item["id"]):
                totals = item["stock_bins_totals"]
                self.assertEqual(totals["count"], expected["kpi3_contracts"])
                self.assertMoney(totals["amount"], expected["kpi3_amount_detail_local"])
                self.assertMoney(sum(row["pct_count"] for row in item["stock_bins"]), 1)
                self.assertMoney(sum(row["pct_amount"] for row in item["stock_bins"]), 1)
                for row in item["stock_bins"]:
                    self.assertMoney(row["pct_count"], row["count"] / totals["count"])
                    self.assertMoney(row["pct_amount"], row["amount"] / totals["amount"])
        italy = country(self.data, "IT")
        self.assertMoney(italy["stock_bins_totals"]["amount"], 787426.24)
        self.assertMoney(italy["metrics"]["identified_eur"], 3127145.95)
        self.assertMoney(italy["metrics"]["identified_eur"] - 3023244.63, 103901.32)
        self.assertIsNone(italy["stock_vintages"][0]["amount"])
        self.assertIsNone(italy["stock_vintages"][0]["pct_amount"])
        self.assertEqual(country(self.data, "ZA")["stock_vintages"][0]["label"], "<2019")

    def test_missing_previous_quarter_is_unknown_not_zero_or_an_older_observation(self):
        self.assertEqual(self.data["previous_period"], "2026Q1")
        self.assertFalse(self.data["previous_available"])
        for item in self.data["countries"]:
            self.assertIsNone(item["previous_metrics"])
            self.assertIsNone(item["metrics"]["stock_change"])
        # Having a historical observation is not evidence for the missing Q1.
        older = period_copy(self.snapshot, "2025Q4")
        history = merge_snapshot(self.history, older)
        data = build_presentation_data(history, "2026Q2", self.config)
        self.assertFalse(data["previous_available"])
        self.assertIsNone(country(data, "ZA")["previous_metrics"])

    def test_previous_quarter_keeps_its_fx_and_missing_entity_stays_unknown(self):
        previous = period_copy(self.snapshot, "2026Q1", rate_multiplier=2)
        previous["entities"] = [item for item in previous["entities"] if item["id"] != "AT"]
        history = merge_snapshot(self.history, previous)
        before_history, before_config = deepcopy(history), deepcopy(self.config)
        data = build_presentation_data(history, "2026Q2", self.config)
        self.assertTrue(data["previous_available"])
        za = country(data, "ZA")
        self.assertMoney(za["previous_metrics"]["identified_eur"], self.expected["South Africa"]["credit_stock_including_unreachable_eur"] / 2)
        self.assertMoney(za["metrics"]["identified_eur"], self.expected["South Africa"]["credit_stock_including_unreachable_eur"])
        self.assertMoney(za["metrics"]["stock_change"], 1)
        self.assertIsNone(country(data, "AT")["previous_metrics"])
        self.assertIsNone(country(data, "AT")["metrics"]["stock_change"])
        self.assertEqual(history, before_history)
        self.assertEqual(self.config, before_config)

    def test_unknown_month_is_not_converted_to_zero_and_zero_denominator_has_no_rate(self):
        snapshot = deepcopy(self.snapshot)
        italy = source_entity(snapshot, "IT")
        italy["flows"][0]["amount"] = None
        italy["flows"][1]["count"] = 0
        data = build_presentation_data(merge_snapshot({"snapshots": {}}, snapshot), "2026Q2", self.config)
        result = country(data, "IT")
        self.assertIsNone(result["monthly"][0]["amount"])
        self.assertIsNone(result["monthly"][0]["amount_eur"])
        self.assertIsNone(result["monthly"][0]["remaining"])
        self.assertIsNone(result["monthly"][0]["efficiency_amount"])
        self.assertIsNone(result["monthly"][1]["efficiency_count"])
        self.assertIsNone(result["metrics"]["flow_eur"])

    def test_global_efficiencies_are_weighted_and_fx_converted_before_aggregation(self):
        countries = self.data["countries"]
        metrics = self.data["global"]["metrics"]
        self.assertMoney(metrics["oldest_amount"],
                         sum(c["monthly"][0]["processed_amount_eur"] for c in countries)
                         / sum(c["monthly"][0]["amount_eur"] for c in countries))
        self.assertMoney(metrics["oldest_count"],
                         sum(c["monthly"][0]["processed_count"] for c in countries)
                         / sum(c["monthly"][0]["count"] for c in countries))
        self.assertMoney(metrics["efficiency_amount"],
                         sum(c["metrics"]["processed_eur"] for c in countries)
                         / sum(c["metrics"]["flow_eur"] for c in countries))
        self.assertNotAlmostEqual(metrics["oldest_amount"],
                                  sum(c["metrics"]["oldest_amount"] for c in countries) / len(countries), places=4)
        self.assertNotAlmostEqual(metrics["oldest_amount"], metrics["efficiency_amount"], places=4)
        self.assertEqual(previous_period("Q1_2027"), "2026Q4")


class DetailedWorkbookTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        snapshot = extract_folder(ROOT / "Q2_2026", ROOT / "config")
        cls.history = merge_snapshot({"snapshots": {}}, snapshot)
        cls.config = json.loads((ROOT / "config/presentation.json").read_text(encoding="utf-8"))
        cls.data = build_presentation_data(cls.history, "2026Q2", cls.config)
        cls.temp = tempfile.TemporaryDirectory()
        cls.addClassCleanup(cls.temp.cleanup)
        cls.path = Path(cls.temp.name) / "detailed.xlsx"
        write_detailed_workbook(cls.path, cls.history, "2026Q2", cls.config)
        cls.formulas = openpyxl.load_workbook(cls.path, data_only=False)
        cls.values = openpyxl.load_workbook(cls.path, data_only=True)
        cls.addClassCleanup(cls.formulas.close)
        cls.addClassCleanup(cls.values.close)

    def test_missing_prior_period_is_not_a_bare_reference_to_a_blank_chart_cell(self):
        # Chart gaps must be blank, but visible financial tables must keep n.d.
        # A direct formula reference to such a blank evaluates to zero in Excel.
        for sheet, count in (("Global", 6), ("CORE Stock", 3), ("NON CORE Stock", 3)):
            for row in range(11, 11 + count):
                for column in ("H", "N"):
                    with self.subTest(sheet=sheet, cell=f"{column}{row}"):
                        formula = self.formulas[sheet][f"{column}{row}"].value
                        self.assertEqual(self.values[sheet][f"{column}{row}"].value, "n.d.")
                        if isinstance(formula, str) and formula.startswith("="):
                            self.assertIn("IF(", formula)
                            self.assertIn('"n.d."', formula)
            self.assertEqual(self.values[sheet][f"H{11 + count}"].value, "n.d.")

    def test_country_charts_have_native_amount_stacks_and_secondary_percentage_lines(self):
        with ZipFile(self.path) as archive:
            self.assertFalse(any(name.startswith("xl/externalLinks/") for name in archive.namelist()))
            native_charts = [ET.fromstring(archive.read(name)) for name in archive.namelist()
                             if name.startswith("xl/charts/chart") and name.endswith(".xml")]
        country_charts = []
        for chart in native_charts:
            title = " ".join(chart.itertext())
            if ": Monthly inflow in " in title:
                country_charts.append(chart)
                bars = chart.find(".//c:barChart", CHART_NS)
                lines = chart.find(".//c:lineChart", CHART_NS)
                self.assertIsNotNone(bars)
                self.assertIsNotNone(lines)
                self.assertEqual(bars.find("c:grouping", CHART_NS).attrib["val"], "stacked")
                self.assertEqual(len(lines.findall("c:ser", CHART_NS)), 2)
                self.assertNotEqual([node.attrib["val"] for node in bars.findall("c:axId", CHART_NS)],
                                    [node.attrib["val"] for node in lines.findall("c:axId", CHART_NS)])
                line_axis_id = lines.findall("c:axId", CHART_NS)[-1].attrib["val"]
                secondary = next(axis for axis in chart.findall(".//c:valAx", CHART_NS)
                                 if axis.find("c:axId", CHART_NS).attrib["val"] == line_axis_id)
                self.assertEqual(secondary.find("c:numFmt", CHART_NS).attrib["formatCode"], "0%")
                self.assertEqual(secondary.find("c:axPos", CHART_NS).attrib["val"], "r")
                self.assertTrue(all("GraphData!" in ref.text or "'GraphData'!" in ref.text
                                    for ref in chart.findall(".//c:f", CHART_NS)))
        self.assertEqual(len(country_charts), len(self.data["countries"]))

    def test_financial_totals_and_formula_caches_remain_consistent(self):
        self.assertAlmostEqual(self.values["Global"]["K17"].value,
                               self.data["global"]["metrics"]["identified_eur"], places=5)
        self.assertEqual(self.values["IT Italy"]["X12"].value, 5595)
        self.assertAlmostEqual(self.values["IT Italy"]["Z12"].value, 787426.24, places=2)
        self.assertEqual(self.values["IT Italy"]["Z14"].value, "n.d.")
        errors = {"#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A", "#NUM!", "#NULL!"}
        for sheet in self.values:
            for row in sheet:
                for cell in row:
                    if cell.data_type == "e" or (isinstance(cell.value, str) and cell.value in errors):
                        self.fail(f"Error cache in {sheet.title}!{cell.coordinate}: {cell.value}")

    def test_large_country_scope_and_comment_overflow_remain_inside_print_layout(self):
        history, config = deepcopy(self.history), deepcopy(self.config)
        snapshot = history["snapshots"]["2026Q2"]
        standard = source_entity(snapshot, "AT")
        for index in range(7):
            entity = deepcopy(standard)
            entity.update(id=f"X{index}", name=f"Additional country {index}")
            snapshot["entities"].append(entity)
        comments = {"GLOBAL": [f"Global business comment {index}" for index in range(10)],
                    "IT": [f"Italy business comment {index}" for index in range(10)]}
        config["comments"]["2026Q2"].update(comments)
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "large_scope.xlsx"
            write_detailed_workbook(path, history, "2026Q2", config)
            workbook = openpyxl.load_workbook(path, data_only=True)
            try:
                for title in ("Global", "CORE Flows", "IT Italy", "Presentation inputs"):
                    sheet = workbook[title]
                    occupied = set()
                    for merged in sheet.merged_cells.ranges:
                        cells = set(merged.cells)
                        self.assertFalse(cells & occupied, f"Overlapping ranges on {title}")
                        occupied.update(cells)
                text = [str(cell.value) for sheet in workbook for row in sheet for cell in row
                        if cell.value is not None]
                for lines in comments.values():
                    for line in lines:
                        self.assertIn(line, text)
                for title in ("Global", "CORE Flows"):
                    sheet = workbook[title]
                    print_end = openpyxl.utils.cell.range_boundaries(str(sheet.print_area).split("!")[-1])[3]
                    # The comparison tables include both O (amount) and X (count)
                    # country names; neither may be cut off by the footer/print area.
                    country_names = {entity["name"] for entity in snapshot["entities"]}
                    comparison_rows = [cell.row for row in sheet for cell in row
                                       if cell.column in (15, 24) and cell.value in country_names]
                    footers = [cell.row for row in sheet for cell in row
                               if isinstance(cell.value, str) and cell.value.startswith("Positive Balance     ")]
                    self.assertTrue(comparison_rows)
                    self.assertGreater(max(footers), max(comparison_rows))
                    self.assertGreaterEqual(print_end, max(footers))
            finally:
                workbook.close()

    def test_cli_both_then_detailed_preserves_existing_summary(self):
        from positive_balance.cli import main

        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            summary = root / "summary.xlsx"
            detailed = root / "detailed.xlsx"
            history = root / "history.json"
            args = [str(ROOT / "Q2_2026"), "--config", str(ROOT / "config"),
                    "--history", str(history)]
            with redirect_stdout(io.StringIO()), redirect_stderr(io.StringIO()):
                code = main(args + ["--report", "both", "--output", str(summary),
                                    "--detailed-output", str(detailed)])
            self.assertEqual(code, 0)
            self.assertTrue(summary.is_file())
            self.assertTrue(detailed.is_file())
            old_summary = hashlib.sha256(summary.read_bytes()).hexdigest()
            with redirect_stdout(io.StringIO()), redirect_stderr(io.StringIO()):
                code = main(args + ["--report", "detailed", "--output", str(detailed)])
            self.assertEqual(code, 0)
            self.assertEqual(hashlib.sha256(summary.read_bytes()).hexdigest(), old_summary)
            state = json.loads(history.read_text(encoding="utf-8"))
            self.assertEqual(list(state["snapshots"]), ["2026Q2"])


if __name__ == "__main__":
    unittest.main()
