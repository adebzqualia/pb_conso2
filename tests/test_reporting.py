"""Regression checks against the supplied reporting examples.

Run with ``python -m unittest discover -s tests -v``. Excel error cases modify
temporary ZIP copies so the original files, and their formula caches, survive.
"""
from __future__ import annotations

from contextlib import redirect_stderr, redirect_stdout
from copy import deepcopy
import hashlib
import io
import json
from pathlib import Path, PurePosixPath
import shutil
import tempfile
import unittest
import xml.etree.ElementTree as ET
from zipfile import ZipFile

import openpyxl

from positive_balance.history import atomic_json, history_lock, load_history, merge_snapshot
from positive_balance.model import complete_sum, parse_period, quarter_end, ratio, rolling_months
from positive_balance.reader import extract_folder


ROOT = Path(__file__).resolve().parents[1]
SAMPLES = ROOT / "Q2_2026"
CONFIG = ROOT / "config"
NS = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
REL_ID = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"


def hashes(folder):
    return {p.name: hashlib.sha256(p.read_bytes()).hexdigest()
            for p in folder.iterdir() if p.is_file()}


def edit_cell(path, sheet, address, *, value=None, remove_cache=False, remove_formula=False):
    """Patch one stored value without openpyxl discarding other formula caches."""
    with ZipFile(path) as source:
        entries = {item.filename: (item, source.read(item.filename))
                   for item in source.infolist()}
    workbook = ET.fromstring(entries["xl/workbook.xml"][1])
    rid = next(node.attrib[REL_ID] for node in workbook.findall("s:sheets/s:sheet", NS)
               if node.attrib["name"] == sheet)
    relationships = ET.fromstring(entries["xl/_rels/workbook.xml.rels"][1])
    target = next(node.attrib["Target"] for node in relationships
                  if node.attrib["Id"] == rid)
    member = target.lstrip("/") if target.startswith("/") else str(PurePosixPath("xl") / target)
    xml = ET.fromstring(entries[member][1])
    cell = xml.find(f".//s:c[@r='{address}']", NS)
    if cell is None:
        raise AssertionError(f"Fixture cell does not exist: {sheet}!{address}")
    old = cell.find("s:v", NS)
    if old is not None:
        cell.remove(old)
    if remove_formula and cell.find("s:f", NS) is not None:
        cell.remove(cell.find("s:f", NS))
    if remove_cache:
        if cell.find("s:f", NS) is None:
            raise AssertionError(f"Expected a formula fixture: {sheet}!{address}")
    elif value is not None:
        ET.SubElement(cell, f"{{{NS['s']}}}v").text = str(value)
    entries[member] = (entries[member][0], ET.tostring(xml, encoding="utf-8", xml_declaration=True))
    with ZipFile(path, "w") as output:
        for info, content in entries.values():
            output.writestr(info, content)


class ReportingTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.before = hashes(SAMPLES)
        cls.snapshot = extract_folder(SAMPLES, CONFIG)
        cls.expected = json.loads((ROOT / "analysis/reconciliation_expected.json").read_text(encoding="utf-8"))

    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix="positive_balance_test_")
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)

    def copied_inputs(self):
        folder = self.root / "Q2_2026"
        config = self.root / "custom_config"
        shutil.copytree(SAMPLES, folder)
        shutil.copytree(CONFIG, config)
        return folder, config

    def assertMoney(self, actual, expected):
        self.assertIsNotNone(actual)
        self.assertAlmostEqual(actual, expected, delta=.00001)

    def test_sources_unchanged_and_six_entities_are_auditable(self):
        self.assertEqual(hashes(SAMPLES), self.before)
        self.assertEqual(self.snapshot["period"], "2026Q2")
        self.assertEqual({e["id"] for e in self.snapshot["entities"]}, {"AT", "IT", "ZA", "DK", "NO", "SE"})
        self.assertEqual(len(self.snapshot["sources"]), 5)  # Four country books plus FX.
        for source in self.snapshot["sources"]:
            self.assertEqual(source["sha256"], self.before[source["file"]])
        self.assertEqual({item["file"] for item in self.snapshot["excluded"]},
                         {p for p in ("ROMANIA.xlsx", "BRAZIL.xlsb") if (SAMPLES / p).exists()})
        self.assertEqual({i["entity"] for i in self.snapshot["issues"] if i["code"] == "CST_REQUIRED"}, {"DK", "NO", "SE"})

    def test_stock_and_six_month_flows_match_independent_source_controls(self):
        rates = {r["currency"]: r for r in self.snapshot["fx"]}
        for entity in self.snapshot["entities"]:
            with self.subTest(country=entity["name"]):
                expected = self.expected[entity["name"]]
                rate = rates[entity["currency"]]
                self.assertMoney(rate["rate"], expected["closing_local_per_eur"])
                self.assertEqual(rate["date"][:7], "2026-06")
                self.assertEqual(rate["quote"], "local_per_EUR")
                generated = sum(r["total"] for r in entity["stock"])
                processed = sum(r[k] for r in entity["stock"] for k in entity["processed_fields"])
                remaining = generated - processed
                unreachable = sum(r.get("unreachable") or 0 for r in entity["stock"])
                flow = sum(r["amount"] for r in entity["flows"])
                treated = sum(r[k] for r in entity["flows"] for k in entity["processed_fields"])
                self.assertMoney(generated, expected["total_generated_local"])
                self.assertMoney(processed, expected["total_processed_local"])
                self.assertMoney(remaining, expected["remaining_local"])
                self.assertMoney((remaining + unreachable) / rate["rate"], expected["credit_stock_including_unreachable_eur"])
                self.assertMoney(flow / rate["rate"], expected["flow_6m_eur"])
                self.assertMoney(treated / rate["rate"], expected["processed_6m_eur"])
                self.assertMoney(sum(r["amount"] for r in entity["stock_bins"]), expected["kpi3_amount_detail_local"])
                self.assertEqual(sum(r["count"] for r in entity["stock_bins"]), expected["kpi3_contracts"])

    def test_semester_quarter_and_oldest_cohort_are_distinct(self):
        italy = next(e for e in self.snapshot["entities"] if e["id"] == "IT")
        expected = self.expected["Italy"]
        flows = italy["flows"]
        self.assertEqual([r["month"] for r in flows], [f"2026-{m:02}" for m in range(1, 7)])
        total = sum(r["amount"] for r in flows)
        treated = sum(r[k] for r in flows for k in italy["processed_fields"])
        weighted = treated / total
        oldest = flows[0]["processed_amount"] / flows[0]["amount"]
        self.assertMoney(weighted, expected["weighted_efficiency_amount_6m"])
        self.assertMoney(oldest, expected["oldest_cohort_efficiency_amount"])
        self.assertGreater(abs(oldest - weighted), .05)
        self.assertNotAlmostEqual(weighted, sum(r["processed_amount"] / r["amount"] for r in flows) / 6, places=5)
        self.assertLess(sum(r["amount"] for r in flows if r["month"] >= "2026-04"), total)
        source_stock = sum(r["total"] - sum(r[k] for k in italy["processed_fields"]) + r["unreachable"] for r in italy["stock"])
        self.assertMoney(source_stock - expected["legacy_credit_stock_including_unreachable_eur"], 103901.32)

    def test_period_boundaries_and_unknown_values(self):
        self.assertEqual(parse_period("Q1_2027"), "2027Q1")
        self.assertEqual(quarter_end("2028Q1").isoformat(), "2028-03-31")
        self.assertEqual(rolling_months("2027Q1"), ["2027-03", "2027-02", "2027-01", "2026-12", "2026-11", "2026-10"])
        with self.assertRaises(ValueError):
            parse_period("Q5_2026")
        self.assertIsNone(complete_sum([100, None]))
        self.assertEqual(complete_sum([100, 0]), 100)
        self.assertIsNone(ratio(1, 0))
        self.assertIsNone(ratio(None, 5))

    def test_history_replaces_one_period_and_keeps_overlapping_observations(self):
        first = deepcopy(self.snapshot)
        # A deliberately synthetic observation demonstrates the history key.
        later = deepcopy(first)
        later["period"] = "2026Q3"
        later["entities"][0]["flows"][0]["processed_amount"] = 123
        state = merge_snapshot({"snapshots": {}}, first)
        state = merge_snapshot(state, later)
        rerun = deepcopy(first)
        rerun["entities"][0]["flows"][0]["processed_amount"] = 456
        state = merge_snapshot(state, rerun)
        self.assertEqual(list(state["snapshots"]), ["2026Q2", "2026Q3"])
        self.assertEqual(state["snapshots"]["2026Q2"]["entities"][0]["flows"][0]["processed_amount"], 456)
        self.assertEqual(state["snapshots"]["2026Q3"]["entities"][0]["flows"][0]["processed_amount"], 123)
        path = self.root / "history.json"
        atomic_json(path, state)
        self.assertEqual(load_history(path), state)
        with history_lock(path):
            with self.assertRaisesRegex(ValueError, "verrou"):
                with history_lock(path):
                    self.fail("Concurrent lock should not have been granted")
        self.assertFalse(path.with_suffix(".json.lock").exists())

    def test_configuration_controls_filenames_layout_and_units(self):
        folder, config = self.copied_inputs()
        renamed = folder / "AT_custom.xlsx"
        (folder / "AUSTRIA.xlsx").rename(renamed)
        countries_path = config / "countries.json"
        countries = json.loads(countries_path.read_text(encoding="utf-8"))
        austria = next(e for e in countries["entities"] if e["id"] == "AT")
        austria.update(file=renamed.name, units=10, profile="custom_standard")
        countries_path.write_text(json.dumps(countries), encoding="utf-8")
        layouts_path = config / "layouts.json"
        layouts = json.loads(layouts_path.read_text(encoding="utf-8"))
        layouts["profiles"]["custom_standard"] = deepcopy(layouts["profiles"]["standard"])
        layouts_path.write_text(json.dumps(layouts), encoding="utf-8")
        custom = extract_folder(folder, config)
        actual = next(e for e in custom["entities"] if e["id"] == "AT")
        original = next(e for e in self.snapshot["entities"] if e["id"] == "AT")
        self.assertMoney(actual["flows"][0]["amount"], original["flows"][0]["amount"] * 10)
        self.assertEqual(actual["flows"][0]["count"], original["flows"][0]["count"])
        self.assertIn("AT_custom.xlsx!", actual["flows"][0]["source"])

    def test_contradictory_folder_reporting_and_fx_periods_fail(self):
        with self.assertRaisesRegex(ValueError, "Folder period"):
            extract_folder(SAMPLES, CONFIG, "2026Q3")
        folder, config = self.copied_inputs()
        edit_cell(folder / "AUSTRIA.xlsx", "KPI_2", "G4", value=46000)
        with self.assertRaisesRegex(ValueError, "reporting date"):
            extract_folder(folder, config)
        shutil.copy2(SAMPLES / "AUSTRIA.xlsx", folder / "AUSTRIA.xlsx")
        edit_cell(folder / "Taux.xlsx", "Table des taux M", "B2", value=46000)
        with self.assertRaisesRegex(ValueError, "FX date"):
            extract_folder(folder, config)

    def test_missing_fx_and_uncalculated_formula_fail(self):
        folder, config = self.copied_inputs()
        fx = openpyxl.load_workbook(folder / "Taux.xlsx", read_only=True, data_only=True)
        try:
            sheet = fx["Table des taux M"]
            zar_row = next(row for row in range(4, sheet.max_row + 1) if sheet.cell(row, 1).value == "ZAR")
        finally:
            fx.close()
        edit_cell(folder / "Taux.xlsx", "Table des taux M", f"E{zar_row}", value=None)
        with self.assertRaisesRegex(ValueError, "Missing or nonpositive FX rate for ZAR"):
            extract_folder(folder, config)
        shutil.copy2(SAMPLES / "Taux.xlsx", folder / "Taux.xlsx")
        edit_cell(folder / "AUSTRIA.xlsx", "KPI_2", "B11", remove_cache=True)
        with self.assertRaisesRegex(ValueError, "Missing formula cache"):
            extract_folder(folder, config)

    def test_unknown_excel_and_duplicate_entity_mapping_fail(self):
        folder, config = self.copied_inputs()
        shutil.copy2(folder / "AUSTRIA.xlsx", folder / "UNMAPPED.xlsx")
        with self.assertRaisesRegex(ValueError, "Unconfigured Excel input UNMAPPED.xlsx"):
            extract_folder(folder, config)
        (folder / "UNMAPPED.xlsx").unlink()
        path = config / "countries.json"
        countries = json.loads(path.read_text(encoding="utf-8"))
        duplicate = dict(countries["entities"][0], id="DUPLICATE")
        countries["entities"].append(duplicate)
        path.write_text(json.dumps(countries), encoding="utf-8")
        with self.assertRaisesRegex(ValueError, "double-count"):
            extract_folder(folder, config)

    def test_missing_numeric_is_preserved_as_unknown(self):
        folder, config = self.copied_inputs()
        edit_cell(folder / "ITALY.xlsx", "KPI_2", "C27", value=None, remove_formula=True)
        snapshot = extract_folder(folder, config)
        italy = next(e for e in snapshot["entities"] if e["id"] == "IT")
        self.assertIsNone(italy["flows"][-1]["amount"])
        self.assertIsNone(complete_sum(r["amount"] for r in italy["flows"]))
        self.assertTrue(any(i["code"] == "MISSING_VALUE" and "C27" in i["source"] for i in snapshot["issues"]))

    def test_output_caches_formulas_tables_and_unknowns(self):
        from positive_balance.workbook import write_workbook, entity_metrics
        output = self.root / "synthesis.xlsx"
        history = merge_snapshot({"snapshots": {}}, self.snapshot)
        write_workbook(output, history, "2026Q2")
        values = openpyxl.load_workbook(output, data_only=True)
        formulas = openpyxl.load_workbook(output, data_only=False)
        try:
            self.assertEqual(values["Synthese"]["C4"].value, "2026Q2")
            self.assertTrue(any("C4" in v.sqref for v in formulas["Synthese"].data_validations.dataValidation))
            self.assertEqual(sum(len(s.tables) for s in formulas), 6)
            self.assertFalse(formulas._external_links)
            for row, entity in enumerate(self.snapshot["entities"], 6):
                expected = self.expected[entity["name"]]
                self.assertMoney(values["Consolidation"][f"M{row}"].value, expected["remaining_eur"])
                self.assertMoney(values["Consolidation"][f"N{row}"].value, expected["credit_stock_including_unreachable_eur"])
                self.assertMoney(values["Consolidation"][f"Q{row}"].value, expected["flow_6m_eur"])
                self.assertMoney(values["Consolidation"][f"U{row}"].value, expected["weighted_efficiency_amount_6m"])
                self.assertMoney(values["Consolidation"][f"X{row}"].value, expected["oldest_cohort_efficiency_amount"])
                self.assertTrue(formulas["Consolidation"][f"M{row}"].value.startswith("=IF("))
                if entity.get("requires_cst"):
                    self.assertEqual(values["Consolidation"][f"O{row}"].value, "n.d.")
            errors=[(s.title,c.coordinate,c.value) for s in values for row in s for c in row if c.data_type=="e"]
            self.assertEqual(errors, [])
            changed = deepcopy(self.snapshot["entities"][0])
            changed["flows"][0]["amount"] = None
            self.assertIsNone(entity_metrics(changed, 1, "2026Q2")["flow_eur"])
        finally:
            values.close()
            formulas.close()

    def test_next_period_history_grows_without_repricing_previous_period(self):
        from positive_balance.workbook import write_workbook
        first = deepcopy(self.snapshot)
        later = deepcopy(first)
        later["period"] = "2026Q3"
        for entity in later["entities"]:
            for record, month in zip(entity["flows"], reversed(rolling_months("2026Q3"))):
                record["month"] = month
        for rate in later["fx"]:
            rate["date"] = "2026-09-01"
            if rate["currency"] == "ZAR":
                rate["rate"] *= 2
        history = merge_snapshot(merge_snapshot({"snapshots": {}}, first), later)
        output = self.root / "two_periods.xlsx"
        write_workbook(output, history, "2026Q3")
        book = openpyxl.load_workbook(output, data_only=True)
        try:
            self.assertEqual(book["Consolidation"].max_row, 17)
            self.assertEqual(book["Evolution flux"]["C6"].value,"2026Q2")
            self.assertEqual(book["Evolution flux"]["D6"].value,"2026Q3")
            self.assertEqual(book["Synthese"]["C4"].value,"2026Q3")
            self.assertMoney(book["Consolidation"]["M8"].value, self.expected["South Africa"]["remaining_eur"])
            self.assertMoney(book["Consolidation"]["M14"].value, self.expected["South Africa"]["remaining_eur"] / 2)
        finally:
            book.close()

    def test_cli_validation_failure_and_source_overwrite_are_non_destructive(self):
        from positive_balance.cli import main
        history = self.root / "history.json"
        output = self.root / "result.xlsx"
        with redirect_stdout(io.StringIO()), redirect_stderr(io.StringIO()):
            result = main([str(SAMPLES),"--validate-only","--history",str(history),"--output",str(output)])
            self.assertEqual(result,0)
            self.assertFalse(history.exists())
            self.assertFalse(output.exists())
            result = main([str(SAMPLES),"--output",str(SAMPLES / "AUSTRIA.xlsx"),"--history",str(history)])
            self.assertEqual(result,2)
            self.assertEqual(hashes(SAMPLES),self.before)
            self.assertFalse(history.exists())
            result = main([str(SAMPLES),"--strict","--history",str(history),"--output",str(output)])
            self.assertEqual(result,2)
            self.assertFalse(output.exists())


if __name__ == "__main__":
    unittest.main()
